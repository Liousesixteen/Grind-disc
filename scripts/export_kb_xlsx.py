#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""全量库 xlsx 导出（多 sheet，带来源与更新时间）"""
import datetime as dt, os, sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DB='cache/kb/trading_kb.sqlite'; OUT='reports/全量投研库_20261026.xlsx'
c=sqlite3.connect(DB, timeout=300)
wb=Workbook(); wb.remove(wb.active)
HDR=Font(bold=True,color='FFFFFF'); FILL=PatternFill('solid',fgColor='2F5597')

def sheet(name, sql, params=(), widths=None, note=''):
    ws=wb.create_sheet(name[:31])
    cur=c.execute(sql, params); cols=[d[0] for d in cur.description]
    ws.append(cols)
    for j,_ in enumerate(cols,1):
        cell=ws.cell(1,j); cell.font=HDR; cell.fill=FILL; cell.alignment=Alignment(horizontal='center')
    n=0
    for row in cur:
        ws.append(list(row)); n+=1
    for j,col in enumerate(cols,1):
        w=(widths or {}).get(col, 12 if len(str(col))<8 else 18)
        ws.column_dimensions[get_column_letter(j)].width=w
    ws.freeze_panes='A2'
    ws2=wb.create_sheet(name[:28]+'_说明') if note else None
    if note:
        ws2['A1']='数据来源'; ws2['B1']=note
        ws2['A2']='更新时间'; ws2['B2']=dt.datetime.now().strftime('%Y-%m-%d %H:%M')
        ws2['A3']='行数'; ws2['B3']=n
    return n

# 1 概览
ws=wb.create_sheet('00_概览')
ws.append(['全量投研库 · 导出说明']); ws['A1'].font=Font(bold=True,size=14)
ws.append([]); ws.append(['数据日','2026-10-26']); ws.append(['导出时间',dt.datetime.now().strftime('%Y-%m-%d %H:%M')])
ws.append(['库文件','cache/kb/trading_kb.sqlite'])
ws.append([]); ws.append(['Sheet','内容','行数'])
tabs=[('01_坐标_链环节',"SELECT chain_id 链, name 环节, n_words 词数 FROM dim_segment ORDER BY chain_id, seg_id",'dim_segment（自建坐标系）'),
      ('02_申万三级',"SELECT sw_l3 三级行业, chain_id 对应链, category 类别, rule 规则 FROM rel_sw_coord ORDER BY chain_id, sw_l3",'Tushare index_classify + 自建映射'),
      ('03_概念_东财',"SELECT concept 概念, category 类目, chain_id 链, seg_id 环节, n_stocks 成分数 FROM dim_concept ORDER BY n_stocks DESC",'东财 F10 + concept_map'),
      ('04_概念_同花顺',"SELECT name 板块, category 类目, chain_id 链, CROSS 说明 FROM rel_ths_coord ORDER BY category, name".replace('CROSS','rule'),'Tushare ths_index（无参 2517）+ 对齐'),
      ('05_题材_开盘啦',"SELECT theme_name 题材, category 类目, chain_id 链, seg_id 环节 FROM rel_kpl_coord ORDER BY theme_name",'Tushare kpl_concept'),
      ('06_产品',"SELECT product 产品, n_stocks 覆盖票数, seg_id 归属环节, map_rule 判定依据 FROM dim_product ORDER BY n_stocks DESC",'F10 主营构成（5,576 只）'),
      ('07_ETF',"SELECT code 代码, name 名称, chain_id 链, index_name 跟踪指数 FROM dim_etf ORDER BY chain_id, code",'Tushare etf_basic'),
      ('08_跨链簇',"SELECT members 成员概念, chains 跨链, n_chains 链数, n_concepts 概念数, n_stocks 股票数 FROM theme_cross_cluster ORDER BY n_concepts DESC",'theme_cross.py（结构层）'),
      ('09_簇趋势',"SELECT name 簇, seq20 二十日序列, today 今日, avg5 五日均, avg15 十五日均, trend 趋势 FROM theme_cluster_trend ORDER BY today DESC",'theme_trend.py'),
      ('10_回测',"SELECT bucket 分组, samples 样本, avg_next 次日均值, win_rate 胜率, avg_excess 超额 FROM bt_theme_result",'theme_backtest.py')]
for name,sql,note in tabs:
    n=sheet(name,sql,note=note)
    ws.append([name,note,n])
for j,w in ((1,18),(2,40),(3,40),(4,12)): ws.column_dimensions[get_column_letter(j)].width=w
os.makedirs('reports',exist_ok=True)
wb.save(OUT)
print('已导出', OUT, round(os.path.getsize(OUT)/1024/1024,2), 'MB')
for n in wb.sheetnames: print('  ', n, wb[n].max_row-1, '行')
